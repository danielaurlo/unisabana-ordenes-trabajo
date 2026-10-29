"""Generate a completed institutional Excel workbook in memory."""

from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from gestor_contratos.config.excel_mapping import (
    CERTIFICATE_NO_CELL,
    CERTIFICATE_SHEET,
    CERTIFICATE_TYPE_CELL,
    CERTIFICATE_YES_CELL,
    DEFAULT_TEMPLATE_PATH,
    DETAIL_FIELD_COLUMNS,
    DETAIL_TABLE_END_ROW,
    DETAIL_TABLE_SHEET,
    DETAIL_TABLE_START_ROW,
    EXCEL_FIELD_MAP,
    ExcelCellTarget,
)
from gestor_contratos.models.orden_trabajo import OrdenTrabajoData


class ExcelGenerationError(Exception):
    """Raised when the workbook cannot be generated safely."""


class ExcelGenerator:
    """Fill the institutional workbook without modifying its source file."""

    def __init__(
        self,
        template_path: Path | str = DEFAULT_TEMPLATE_PATH,
        field_map: dict[str, tuple[ExcelCellTarget, ...]] = EXCEL_FIELD_MAP,
    ) -> None:
        self.template_path = Path(template_path)
        self.field_map = field_map

    def generate(self, data: OrdenTrabajoData) -> bytes:
        """Return a completed workbook as bytes."""
        if data.campos_faltantes or data.campos_ambiguos:
            raise ExcelGenerationError(
                "No se puede generar el Excel: hay campos faltantes o ambiguos."
            )
        if not self.template_path.is_file():
            raise ExcelGenerationError(
                f"No se encontró la plantilla Excel: {self.template_path}"
            )

        try:
            workbook = load_workbook(self.template_path, data_only=False)
            for field_path, targets in self.field_map.items():
                value = self._resolve_field(data, field_path)
                for target in targets:
                    self._write_value(workbook, target, value)
            self._write_details(workbook, data)
            self._write_certificate_requirement(workbook, data)

            output = BytesIO()
            workbook.save(output)
            return output.getvalue()
        except ExcelGenerationError:
            raise
        except (KeyError, OSError, ValueError) as exc:
            raise ExcelGenerationError(
                "No fue posible completar la plantilla Excel."
            ) from exc

    @staticmethod
    def _resolve_field(data: OrdenTrabajoData, field_path: str) -> Any:
        value: Any = data
        for part in field_path.split("."):
            value = getattr(value, part)
        return value

    @staticmethod
    def _write_value(workbook: Any, target: ExcelCellTarget, value: Any) -> None:
        if value is None:
            return
        if target.sheet not in workbook.sheetnames:
            raise ExcelGenerationError(
                f"No existe la hoja configurada: {target.sheet}"
            )
        workbook[target.sheet][target.cell] = value

    @staticmethod
    def _write_details(workbook: Any, data: OrdenTrabajoData) -> None:
        if len(data.detalles) > DETAIL_TABLE_END_ROW - DETAIL_TABLE_START_ROW + 1:
            raise ExcelGenerationError(
                "La plantilla no tiene suficientes filas para todos los detalles."
            )
        if DETAIL_TABLE_SHEET not in workbook.sheetnames:
            raise ExcelGenerationError(
                f"No existe la hoja configurada: {DETAIL_TABLE_SHEET}"
            )

        sheet = workbook[DETAIL_TABLE_SHEET]
        for offset, detail in enumerate(data.detalles):
            row = DETAIL_TABLE_START_ROW + offset
            for field_name, column in DETAIL_FIELD_COLUMNS.items():
                ExcelGenerator._write_value(
                    workbook,
                    ExcelCellTarget(DETAIL_TABLE_SHEET, f"{column}{row}"),
                    getattr(detail, field_name),
                )
            if detail.numero is not None:
                sheet[f"B{row}"] = detail.numero

    @staticmethod
    def _write_certificate_requirement(
        workbook: Any, data: OrdenTrabajoData
    ) -> None:
        """Mark the certificate choice and its description on the acta."""
        if CERTIFICATE_SHEET not in workbook.sheetnames:
            raise ExcelGenerationError(
                f"No existe la hoja configurada: {CERTIFICATE_SHEET}"
            )

        sheet = workbook[CERTIFICATE_SHEET]
        requires_certificate = getattr(
            data, "requiere_certificado", False
        )
        certificate_type = getattr(data, "tipo_certificado", None)
        sheet[CERTIFICATE_NO_CELL] = (
            None if requires_certificate else "X"
        )
        sheet[CERTIFICATE_YES_CELL] = (
            "X" if requires_certificate else None
        )
        sheet[CERTIFICATE_TYPE_CELL] = (
            certificate_type if requires_certificate else None
        )
