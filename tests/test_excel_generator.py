from datetime import date
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from gestor_contratos.models.orden_trabajo import (
    DetalleOrden,
    OrdenTrabajoData,
    ProveedorData,
)
from gestor_contratos.services.excel_generator import ExcelGenerator

TEMPLATE = Path("assets/plantilla/3A GROUP SAS-1.xlsx")
SHEET = "Acta de Entrega Final de Manten"


def sample_data() -> OrdenTrabajoData:
    return OrdenTrabajoData(
        fecha_orden=date(2026, 9, 15),
        fecha_acta=date(2026, 9, 15),
        fecha_vencimiento=date(2026, 9, 30),
        fecha_inicio=date(2026, 9, 15),
        numero_orden="27012",
        numero_orden_original="0000027012",
        tipo_orden="Orden de Trabajo",
        objeto="Tala árbol torcido por riesgo a caída ubicado en puente madera",
        sede="PARQUEADEROS",
        proveedor=ProveedorData(
            nombre="ARBOLES Y PLANTAS LA MANA SAS",
            nit="900983709",
            direccion="VDA FONQUETA LT 3 UN VILLA CONCHA",
            persona_contacto="WILSON RUBEN LOPEZ AREVALO",
            ciudad="CHIA",
            telefono="235-0030",
            correo="WL-MANA@HOTMAIL.COM",
        ),
        detalles=[
            DetalleOrden(
                numero=1,
                descripcion="Servicio de tala",
                unidad="UND",
                cantidad=Decimal("1"),
                valor_unitario=Decimal("650000"),
                valor_total=Decimal("650000"),
            )
        ],
    )


def test_fills_institutional_template_without_changing_source() -> None:
    original_bytes = TEMPLATE.read_bytes()
    original_hash = sha256(original_bytes).hexdigest()
    original_workbook = load_workbook(TEMPLATE, data_only=False)
    original_sheet = original_workbook[SHEET]
    original_merged = set(str(item) for item in original_sheet.merged_cells.ranges)
    original_style = original_sheet["AQ14"]._style

    generated = ExcelGenerator().generate(sample_data())

    result_workbook = load_workbook(BytesIO(generated))
    result_sheet = result_workbook[SHEET]
    assert result_sheet["N11"].value == "ARBOLES Y PLANTAS LA MANA SAS"
    assert result_sheet["G12"].value == "900983709"
    assert result_sheet["M14"].value == "27012"
    assert result_sheet["AJ22"].value == "27012"
    assert result_sheet["AQ14"].value.date() == date(2026, 9, 15)
    assert result_sheet["Y15"].value.date() == date(2026, 9, 30)
    assert result_sheet["J7"].value == 15
    assert result_sheet["L7"].value == 9
    assert result_sheet["N7"].value == 2026
    assert result_sheet["F21"].value == 15
    assert result_sheet["N21"].value == 9
    assert result_sheet["Y21"].value == 2026
    assert (
        set(str(item) for item in result_sheet.merged_cells.ranges) == original_merged
    )
    assert result_sheet["AQ14"]._style == original_style
    order_sheet = result_workbook["Orden de Trabajo-Servicio"]
    assert order_sheet["L7"].value == 15
    assert order_sheet["N7"].value == 9
    assert order_sheet["P7"].value == 2026
    assert order_sheet["AO7"].value == "Orden de Trabajo"
    assert order_sheet["AS7"].value == "27012"
    assert order_sheet["O10"].value == sample_data().objeto
    assert order_sheet["K12"].value == "PARQUEADEROS"
    assert order_sheet["H17"].value == "ARBOLES Y PLANTAS LA MANA SAS"
    assert order_sheet["AP17"].value == "900983709"
    assert order_sheet["I18"].value == "VDA FONQUETA LT 3 UN VILLA CONCHA"
    assert order_sheet["AM18"].value == "WILSON RUBEN LOPEZ AREVALO"
    assert order_sheet["D66"].value == "Servicio de tala"
    assert order_sheet["AE66"].value == "UND"
    assert order_sheet["AK66"].value == Decimal("1")
    assert order_sheet["AP66"].value == Decimal("650000")
    assert order_sheet["AV66"].value == '=IF(AK66="","",AK66*AP66)'
    assert sha256(TEMPLATE.read_bytes()).hexdigest() == original_hash
